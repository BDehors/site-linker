#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 tests/article.html 生成两份文档样例：

  1) tests/samples/article-sample.docx  —— Word 版（python-docx，含超链接字段）
  2) tests/samples/article-sample.pdf   —— PDF 版（Edge/Chromium 无头打印，再把正文
     站点名补成链接注释，等价于"文档自带的出处链接 + 正文可点"）

依赖：Python 3 的 lxml、python-docx、PyMuPDF；PDF 生成需要本机安装 Edge 或 Chromium。
用法：python tools/make-samples.py
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

try:
    from lxml import html as lxml_html
    import docx
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    import pymupdf
except ImportError as exc:  # pragma: no cover
    sys.exit(f"缺少依赖：{exc}。请先安装：pip install lxml python-docx pymupdf")

ROOT = Path(__file__).resolve().parent.parent
ARTICLE = ROOT / "tests" / "article.html"
OUTDIR = ROOT / "tests" / "samples"

# 正文里应变成链接的站点名 -> 官网（与词典口径一致，长词在前，避免嵌套命中）
TERMS = [
    ("国家哲学社会科学文献中心", "https://www.ncpssd.cn"),
    ("新英格兰医学杂志", "https://www.nejm.org"),
    ("World Bank Open Data", "https://data.worldbank.org"),
    ("Our World in Data", "https://ourworldindata.org"),
    ("国家统计局", "https://www.stats.gov.cn"),
    ("国家数据", "https://data.stats.gov.cn"),
    ("中国知网", "https://www.cnki.net"),
    ("万方数据", "https://www.wanfangdata.com.cn"),
    ("世界卫生组织", "https://www.who.int"),
    ("PubMed", "https://pubmed.ncbi.nlm.nih.gov"),
    ("bioRxiv", "https://www.biorxiv.org"),
    ("medRxiv", "https://www.medrxiv.org"),
    ("PubScholar", "https://pubscholar.cn"),
    ("The Lancet", "https://www.thelancet.com"),
    ("人民日报", "http://www.people.com.cn"),
    ("Nature", "https://www.nature.com"),
    ("JAMA", "https://jamanetwork.com"),
    ("DOAJ", "https://doaj.org"),
]
TERM_MAP = dict(TERMS)
TERM_RX = re.compile("(" + "|".join(re.escape(t) for t, _ in TERMS) + ")")

DOC_FONT_CN = "SimSun"
DOC_FONT_EN = "Times New Roman"
HEAD_FONT_CN = "SimHei"
LINK_COLOR = "0B4EA2"


# ---------------------------------------------------------------- HTML 解析

def parse_article():
    """把 article.html 解析成按顺序排列的块列表。"""
    tree = lxml_html.parse(str(ARTICLE))
    wrap = tree.xpath('//article[@id="article"]//div[@class="wrap"]')[0]
    blocks = []
    for el in wrap:
        cls = el.get("class") or ""
        tag = el.tag
        if tag == "span" and "demo-chip" in cls:
            blocks.append(("chip", el))
        elif tag == "h1":
            blocks.append(("title", el))
        elif tag == "p" and "byline" in cls:
            blocks.append(("byline", el))
        elif tag == "p" and "dek" in cls:
            blocks.append(("dek", el))
        elif tag == "hr":
            continue
        elif tag == "h2":
            blocks.append(("h2", el))
        elif tag == "p":
            blocks.append(("p", el))
        elif tag == "blockquote":
            blocks.append(("quote", el))
        elif tag == "div" and "chain" in cls:
            blocks.append(("chain", el))
        elif tag == "ol" and "refs" in cls:
            blocks.append(("refs", el))
        elif tag == "div" and "appendix" in cls:
            blocks.append(("appendix", el))
        elif tag == "div" and "more" in cls:
            continue  # 交互区不进文档
        elif tag == "div" and "article-footer" in cls:
            blocks.append(("note", el))
    return blocks


def segs_of(el, linkify=True):
    """把元素拆成 [('text', s) | ('link', text, href)] 片段，可选把站点名转成链接。"""
    out = []

    def emit_text(s):
        if not s:
            return
        if linkify:
            last = 0
            for m in TERM_RX.finditer(s):
                if m.start() > last:
                    out.append(("text", s[last:m.start()]))
                out.append(("link", m.group(0), TERM_MAP[m.group(0)]))
                last = m.end()
            if last < len(s):
                out.append(("text", s[last:]))
        else:
            out.append(("text", s))

    def walk(node):
        if not isinstance(node.tag, str):
            return
        if node.text:
            emit_text(node.text)
        for ch in node:
            if ch.tag == "a":
                href = ch.get("href")
                txt = "".join(ch.itertext())
                if href and txt:
                    out.append(("link", txt, href))
                if ch.tail:
                    emit_text(ch.tail)
            else:
                walk(ch)
                if ch.tail:
                    emit_text(ch.tail)

    walk(el)
    return out


def plain_text(el):
    return " ".join("".join(el.itertext()).split())


# ---------------------------------------------------------------- DOCX

def _set_fonts(rpr, cn, en):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)


def _style_run(run, cn=DOC_FONT_CN, en=DOC_FONT_EN, size=12.0, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    _set_fonts(rpr, cn, en)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_hyperlink(paragraph, text, url, size, cn=DOC_FONT_CN):
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    _set_fonts(rpr, cn, DOC_FONT_EN)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    for tag in ("w:sz", "w:szCs"):
        sz = OxmlElement(tag)
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)


def _para(doc, segs, size=12.0, indent=True, left=None, cn=DOC_FONT_CN,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, line=1.3, after=6.0, before=0.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = line
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if indent:
        pf.first_line_indent = Pt(24)  # 480 twips = 2 字符
    if left is not None:
        pf.left_indent = Pt(left)
    for seg in segs:
        if seg[0] == "text":
            _style_run(p.add_run(seg[1]), cn=cn, size=size, bold=bold)
        else:
            _add_hyperlink(p, seg[1], seg[2], size, cn=cn)
    return p


def _h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    run = p.add_run(text)
    _style_run(run, cn=HEAD_FONT_CN, en="Arial", size=15.0, bold=True, color="000000")
    pf = p.paragraph_format
    pf.space_before = Pt(16)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    return p


def _footer_page_number(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _style_run(run, size=9.0)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)


def build_docx(blocks, out_path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    _set_fonts(rpr, DOC_FONT_CN, DOC_FONT_EN)
    npf = normal.paragraph_format
    npf.line_spacing = 1.3
    npf.space_after = Pt(6)
    npf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for kind, el in blocks:
        if kind == "title":
            _para(doc, [("text", plain_text(el))], size=16, cn=HEAD_FONT_CN,
                  indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=8)
        elif kind == "chip":
            _para(doc, [("text", "（示例文章：情节与数据均为虚构，仅用于演示；文中出处均为可点击链接）")],
                  size=9, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        elif kind == "byline":
            _para(doc, [("text", plain_text(el))], size=10.5, indent=False,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        elif kind == "dek":
            _para(doc, segs_of(el), size=12, indent=False)
        elif kind == "h2":
            _h2(doc, plain_text(el))
        elif kind == "p":
            _para(doc, segs_of(el))
        elif kind == "quote":
            quote_text = plain_text(el)
            att = el.find("footer")
            _para(doc, [("text", quote_text)], size=12, indent=False, left=24, after=2)
            if att is not None:
                _para(doc, [("text", plain_text(att))], size=10.5, indent=False,
                      left=24, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
        elif kind == "chain":
            title = el.find("b")
            if title is not None:
                _para(doc, [("text", plain_text(title))], size=12, cn=HEAD_FONT_CN,
                      indent=False, bold=True, after=2)
            for i, li in enumerate(el.findall(".//li"), 1):
                _para(doc, [("text", f"{i}. ")] + segs_of(li), size=12, indent=False,
                      left=24, after=0)
            _para(doc, [("text", "")], size=12, indent=False, after=4)
        elif kind == "refs":
            for i, li in enumerate(el.findall("li"), 1):
                _para(doc, [("text", f"[{i}] ")] + segs_of(li), size=10.5, indent=False, after=2)
        elif kind == "appendix":
            head = el.find("h2")
            if head is not None:
                _h2(doc, plain_text(head))
            code = el.find(".//pre/code")
            if code is not None:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing = 1.15
                pf.space_after = Pt(4)
                run = p.add_run("".join(code.itertext()).strip("\n"))
                _style_run(run, cn=DOC_FONT_CN, en="Consolas", size=9.5, color="333333")
            hint = el.find("p")
            if hint is not None:
                _para(doc, segs_of(hint), size=10.5, indent=False, after=8)
        elif kind == "note":
            p = el.find("p")
            if p is not None:
                _para(doc, segs_of(p), size=10.5, indent=False, after=0)

    _footer_page_number(doc)

    cp = doc.core_properties
    cp.title = "睡不够，真的更容易感冒吗：一次从论文到数据的核查（示例）"
    cp.author = "站点链接助手 · 测试套件"
    cp.subject = "示例文章：文章阅读与出处核查（文档样例）"
    cp.comments = "由 tools/make-samples.py 从 tests/article.html 生成"

    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------- PDF

EDGE_CANDIDATES = [
    os.environ.get("EDGE_PATH", ""),
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Edge\Application\msedge.exe"),
    "msedge",
    "chromium",
    "chrome",
]


def find_browser():
    for cand in EDGE_CANDIDATES:
        if not cand:
            continue
        found = shutil.which(cand) if not os.path.isabs(cand) else (cand if Path(cand).exists() else None)
        if found:
            return found
    return None


def _page_code_block_rect(page):
    """附录 BibTeX 代码块区域：其中的站点名不做链接（与页内提示一致）。"""
    top = page.search_for("@article{demo2026sleep")
    bot = page.search_for("虚构条目，仅用于演示")
    rects = []
    if top and bot:
        rects.append(pymupdf.Rect(0, top[0].y0 - 2, page.rect.width, bot[-1].y1 + 2))
    elif top:
        rects.append(pymupdf.Rect(0, top[0].y0 - 2, page.rect.width, page.rect.height))
    elif bot:
        rects.append(pymupdf.Rect(0, 0, page.rect.width, bot[-1].y1 + 2))
    return rects


def build_pdf(out_path):
    browser = find_browser()
    if not browser:
        print("!! 未找到 Edge/Chromium，跳过 PDF 生成")
        return None

    tmpdir = Path(tempfile.mkdtemp(prefix="site-linker-pdf-"))
    tmp_pdf = tmpdir / "article.pdf"
    profile = tmpdir / "profile"
    try:
        subprocess.run(
            [browser, "--headless=new", "--disable-gpu", "--no-first-run",
             "--no-default-browser-check", "--disable-sync", "--no-pdf-header-footer",
             f"--user-data-dir={profile}", f"--print-to-pdf={tmp_pdf}", ARTICLE.as_uri()],
            capture_output=True, timeout=120,
        )
        for _ in range(80):  # Edge 启动器先返回、文件稍后落盘
            if tmp_pdf.exists() and tmp_pdf.stat().st_size > 0:
                time.sleep(0.4)
                break
            time.sleep(0.5)
        if not tmp_pdf.exists() or tmp_pdf.stat().st_size == 0:
            raise RuntimeError("PDF 未生成：请确认 Edge/Chromium 可无头运行")

        doc = pymupdf.open(tmp_pdf)
        doc.set_metadata({
            "title": "睡不够，真的更容易感冒吗：一次从论文到数据的核查（示例）",
            "author": "站点链接助手 · 测试套件",
            "subject": "示例文章：文章阅读与出处核查（文档样例）",
            "creator": "站点链接助手 · 测试套件",
            "producer": "站点链接助手 · 测试套件",
        })

        stats = {t: 0 for t, _ in TERMS}

        # 正文站点名 -> 链接注释（已有的出处链接保留不动，代码块区域跳过）
        for page in doc:
            existing = [l["from"] for l in page.get_links()]
            blocked = _page_code_block_rect(page)
            added = []
            for term, url in TERMS:
                for rect in page.search_for(term):
                    covered = any(rect.intersects(r) for r in existing + added + blocked)
                    if covered:
                        continue
                    page.insert_link({"kind": pymupdf.LINK_URI, "from": rect, "uri": url})
                    added.append(rect)
                    stats[term] += 1

        # 页脚页码（居中裸数字，与 docx 版一致）
        for i, page in enumerate(doc, 1):
            label = str(i)
            tw = pymupdf.get_text_length(label, fontname="helv", fontsize=9)
            page.insert_text((page.rect.width / 2 - tw / 2, page.rect.height - 26),
                             label, fontname="helv", fontsize=9, color=(0.45, 0.45, 0.45))

        doc.save(str(out_path), deflate=True, garbage=3)
        doc.close()
        return stats
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


# ---------------------------------------------------------------- 校验与主流程

def verify_docx(path):
    import zipfile

    z = zipfile.ZipFile(path)
    xml = z.read("word/document.xml").decode("utf-8")
    rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    n_links = xml.count("<w:hyperlink")
    urls = re.findall(r'Target="(https?://[^"]+)"', rels)
    ok = n_links >= 20 and all(k in rels for k in ("cnki.net", "jamanetwork.com", "pubmed", "stats.gov.cn"))
    print(f"  docx: 超链接 {n_links} 个（{len(set(urls))} 个不同 URL）{' ✔' if ok else ' ✘'}")
    return ok


def verify_pdf(path):
    doc = pymupdf.open(path)
    total = 0
    uris = set()
    for page in doc:
        for link in page.get_links():
            if link.get("uri"):
                total += 1
                uris.add(link["uri"])
    text = "".join(p.get_text() for p in doc)
    footer_ok = all(
        any(w[4] == str(i) and w[3] > p.rect.height - 40 for w in p.get_text("words"))
        for i, p in enumerate(doc, 1)
    )
    checks = {
        "页数≥2": len(doc) >= 2,
        "标题在文内": "核查" in text,
        "参考资料在文内": "参考资料" in text,
        "页脚页码": footer_ok,
        "交互区不出现在打印版": "加载更多" not in text,
        "链接≥25": total >= 25,
        "含知网链接": any("cnki" in u for u in uris),
        "含统计局链接": any("stats.gov.cn" in u for u in uris),
    }
    for name, ok in checks.items():
        print(f"  pdf: {name} {'✔' if ok else '✘'}（页数 {len(doc)}，链接 {total}）")
    doc.close()
    return all(checks.values())


def main():
    if not ARTICLE.exists():
        sys.exit(f"找不到 {ARTICLE}")
    OUTDIR.mkdir(parents=True, exist_ok=True)
    blocks = parse_article()
    print(f"解析 article.html：{len(blocks)} 个内容块")

    docx_path = OUTDIR / "article-sample.docx"
    build_docx(blocks, docx_path)
    print(f"已生成 {docx_path.relative_to(ROOT)}")

    print("生成 PDF（无头浏览器打印 + 链接补全）……")
    stats = build_pdf(OUTDIR / "article-sample.pdf")
    if stats is not None:
        pdf_path = OUTDIR / "article-sample.pdf"
        print(f"已生成 {pdf_path.relative_to(ROOT)}")
        missed = [t for t, n in stats.items() if n == 0]
        print(f"  链接补全：共 {sum(stats.values())} 处" + (f"，未命中（可能只在已有链接里）：{missed}" if missed else ""))

    print("校验：")
    ok_docx = verify_docx(docx_path)
    ok_pdf = verify_pdf(OUTDIR / "article-sample.pdf") if stats is not None else True
    if ok_docx and ok_pdf:
        print("全部生成并校验通过。")
    else:
        sys.exit("校验未通过，请检查上面的 ✘ 项")


if __name__ == "__main__":
    main()